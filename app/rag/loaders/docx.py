"""DOCX loader：python-docx 抽段落 + 表格，标题层级写进 metadata。

**DOCX 没有"页"，所以 `page=None`。** 这不是偷懒，是诚实：
Word 的"页"是**渲染结果**，同一份文件在不同打印机/字号下页数不同，
文件里根本没存页码。硬编一个 `page=1` 会让"第 3 页"变成谎言 ——
这与 PDF 那边"页码必须真实"是同一条原则的两面。

## 标题层级为什么要进 metadata

`python-docx` 的 `paragraph.style.name` 会给出 `Heading 1` / `Heading 2` / `Title`。
不抓它，下游只知道"这是一段文字"，不知道"这是第二节的标题"：
- 切分时无法"遇到标题就断开"，只能按字数硬切 → 条款被劈开
- 引用时无法显示"第 2 节 付款方式"，只能显示文件名

**只在 metadata 里放级别（int），不放标题路径**：路径要"回忆上一级标题"，
那是组装逻辑，放在 loader 里就会和切分策略耦合。

## 表格和段落必须分开成不同的 Document

`block="table"` 让切分器把它们单独成块。混进文字流的后果与 PDF 那边一样：
表格被切碎后行列关系没了，"这张表的第三列是什么"就答不出来。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from app.rag import Document, LoadResult
from app.rag.loaders.clean import normalize_text

logger = logging.getLogger(__name__)

# `Heading 1` / `标题 1`（中文版 Word）/ `Title`
_HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*([1-6])$", re.IGNORECASE)
_TITLE_STYLE = re.compile(r"^(?:Title|标题)$", re.IGNORECASE)


def heading_level(style_name: str) -> int | None:
    """把 Word 的样式名翻译成标题级别（1-6）；不是标题返回 `None`。

    **中文版 Word 的样式名是"标题 1"而不是"Heading 1"** —— 只匹配英文名的实现
    在中文字档上会静默失效：所有标题都被当成正文段，而且没有任何报错。
    这类"看起来跑通了"的 bug，比崩溃难查得多。
    """
    if not style_name:
        return None
    name = style_name.strip()
    match = _HEADING_STYLE.match(name)
    if match:
        return int(match.group(1))
    if _TITLE_STYLE.match(name):
        return 1  # 文档大标题按 1 级处理，避免它下面没有层级
    return None


def _table_to_markdown(table: object) -> str:
    """把 docx 表格转成 Markdown 表格（与 PDF 那边保持一致）。

    **与 PDF 共用同一个输出格式**是刻意的：下游（切分、embedding、生成）
    只需要认识一种表格表示。两个 loader 各自发明格式，最终会有一处漏改。

    第一行当表头，分隔行 `| --- |` 是 Markdown 语法要求 —— 少了它，
    有些渲染器会把整张表当普通文本，模型也少一个"这是表头"的强信号。
    """
    rows: list[list[str]] = []
    for row in table.rows:  # type: ignore[attr-defined]
        cells = [normalize_text(cell.text).replace("\n", " ").strip() for cell in row.cells]
        # 全空行没有信息量（合并单元格会产生大量空行），跳过
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""

    width = max(len(r) for r in rows)
    lines: list[str] = []
    for index, row in enumerate(rows):
        padded = row + [""] * (width - len(row))
        # `|` 在单元格里会破坏表格语法，转成全角
        lines.append("| " + " | ".join(c.replace("|", "／") for c in padded) + " |")
        if index == 0:
            lines.append("| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines)


def load_docx(path: str | Path) -> LoadResult:
    """加载 DOCX。

    输出策略：**每个非空段落一个 `Document`**，每张表格一个 `Document`。

    **为什么一段一个 Document 而不是整篇一个**：
    下游要按 token 预算切分。loader 已经给出"段"这个天然边界，
    切分就能尽量不破坏段落；整篇一个的话，切分器只能按字数硬切，
    容易把一句话砍成两半。
    """
    from docx import Document as DocxDocument

    file_path = Path(path)
    result = LoadResult()

    try:
        doc = DocxDocument(str(file_path))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"无法打开 DOCX {file_path.name}：{type(exc).__name__}: {exc}") from exc

    paragraph_count = 0
    heading_count = 0
    for para in doc.paragraphs:
        text = normalize_text(para.text or "")
        if not text:
            continue
        level = heading_level(para.style.name if para.style is not None else "")
        metadata: dict[str, str | int | float | bool] = {
            "block": "heading" if level is not None else "text"
        }
        if level is not None:
            metadata["heading_level"] = level
            heading_count += 1
        paragraph_count += 1
        # `page=None`：见模块 docstring，Word 文件里没有页
        result.documents.append(
            Document(source=file_path.name, page=None, text=text, metadata=metadata)
        )

    table_count = 0
    for table in doc.tables:
        markdown = _table_to_markdown(table)
        if not markdown:
            continue
        table_count += 1
        result.documents.append(
            Document(
                source=file_path.name,
                page=None,
                text=markdown,
                metadata={"block": "table", "page_chars": len(markdown)},
            )
        )

    if not result.documents:
        result.warnings.append(f"{file_path.name} 没有可提取的文本或表格 —— 可能是空文档或纯图片")
    else:
        logger.info(
            "DOCX 加载完成：%s（%d 段 / %d 标题 + %d 表格）",
            file_path.name,
            paragraph_count,
            heading_count,
            table_count,
        )

    return result


__all__ = ["heading_level", "load_docx"]
