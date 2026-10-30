"""Loader 测试：三格式统一解析 + 页码对齐 + 清洗的边界。

**fixture 是脚本生成的**（`tests/fixtures/make_fixtures.py`），不是二进制塞进仓库。
好处：review 时能读到"这份 50 页 PDF 每页有什么"，而不是面对一个无法审查的二进制。

## 这个文件里最重要的两组断言

### 一、页码必须"对得上"，而不是"存在"

只断言"有 50 个 Document"是不够的 —— **页码全错也能过**。
所以断言的是：第 N 条 Document 的正文里必须出现"第 N 章"，
且它的 `page == N`、`N in pages`。这才是"页码对得上"的实质证明。

跨页拼接之后这条更严格：`sample_crosspage.pdf` 里 p1 的段落横跨了 p1 和 p2，
所以它的 `pages` 必须是 `(1, 2)` 而 `page` 是 1 —— **只记起始页是半个真话**。

### 二、清洗不许改写正文

引用里的摘录会被用户拿去搜原文。**动一个字就搜不到**。
所以 `test_cleaning_never_rewrites_body_text` 断言：
清洗后文本（去空白）是清洗前的**子序列**。这条一旦被破坏，
任何"顺手优化一下措辞"的改动都会立刻红。
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from app.rag import Document, LoadResult
from app.rag.loaders import LOADERS, SUPPORTED_EXTENSIONS, load_document
from app.rag.loaders.clean import (
    can_join,
    looks_unfinished,
    merge_wrapped_lines,
    normalize_text,
    strip_page_numbers,
    strip_repeated_lines,
)
from app.rag.loaders.docx import heading_level
from app.rag.loaders.pdf import (
    SCAN_PAGE_CHARS,
    load_pdf,
    table_to_markdown,
)

FIXTURES = Path(__file__).parent / "fixtures"


def _body(result: LoadResult) -> str:
    return "\n".join(d.text for d in result.documents)


def _squash(text: str) -> str:
    """去掉所有空白后比对：空白是排版，字才是内容。"""
    return re.sub(r"\s+", "", text)


# ============================================================ 1. 统一模型


def test_document_rejects_bad_page() -> None:
    """页码必须 >= 1 —— 0 或负数是"从 0 开始"这类混淆的信号，早拦。"""
    with pytest.raises(ValueError, match="page 必须"):
        Document(source="x.pdf", page=0, text="t")


def test_document_location_is_human_readable() -> None:
    assert Document(source="a.pdf", page=3, text="x").location == "a.pdf:p3"
    assert Document(source="a.md", page=None, text="x").location == "a.md"


def test_document_block_defaults_to_text() -> None:
    """没写 `block` 的 Document 视为正文 —— 老代码不用改就能继续用。"""
    doc = Document(source="a.md", page=None, text="x")
    assert doc.block == "text"
    assert doc.is_table is False
    table = Document(source="a.md", page=None, text="| a |", metadata={"block": "table"})
    assert table.is_table is True


def test_load_result_summary() -> None:
    result = LoadResult(
        documents=[
            Document(source="a", page=1, text="12345"),
            Document(source="a", page=2, text="678"),
        ],
        warnings=["某页无文本"],
    )
    assert result.total_chars == 8
    assert result.page_count == 2
    assert result.summary()["warnings"] == 1


# ============================================================ 2. 清洗：不许改写正文


def test_normalize_collapses_whitespace_but_keeps_paragraphs() -> None:
    """压缩多余空格与空行，但**保留段落边界** —— 全压成一行会让摘要串味。"""
    assert normalize_text("第一段   有  多余空格\n\n\n\n第二段") == "第一段 有 多余空格\n\n第二段"


def test_normalize_handles_fullwidth_space_and_crlf() -> None:
    assert normalize_text("a\u3000b\r\nc") == "a b\nc"


def test_normalize_empty() -> None:
    assert normalize_text("") == ""
    assert normalize_text("   \n\t  ") == ""


def test_cleaning_never_rewrites_body_text() -> None:
    """**这一段是本文件的核心不变量**：清洗后的正文必须是清洗前的字符子序列。

    为什么这条这么重要：前端引用里给的是原文摘录，用户会拿去 PDF 里搜索。
    清洗只要动了一个字（换词、改标点、调顺序），用户就搜不到 ——
    那一刻他不是在怀疑代码，是在怀疑整份报告。

    唯一允许的删除是**整行被判为噪声**（页眉/页脚/页码），那属于"去掉不该在的东西"，
    而不是"改写还在的东西"。

    **表格块单独判**：表格我们是重写成 Markdown 的（补 `| --- |` 分隔行、
    把单元格里的 `|` 转义），它天然不是原文字序列。对表格的要求换成
    "所有非空单元格都能在原文里找到" —— 那才是表格版"没改写内容"。
    """
    for name in ("sample50.pdf", "sample_wrapped.pdf", "sample.docx", "sample.md"):
        raw = _raw_text(FIXTURES / name)
        result = load_document(FIXTURES / name)

        body = _squash("\n".join(d.text for d in result.documents if not d.is_table))
        assert _is_subsequence(body, _squash(raw)), f"{name} 的正文清洗改写了内容"

        raw_squashed = _squash(raw)
        for doc in result.documents:
            if not doc.is_table:
                continue
            for cell in _table_cells(doc.text):
                assert cell.replace(" ", "") in raw_squashed, (
                    f"{name} 的表格单元格 {cell!r} 在原文里找不到"
                )


def _table_cells(markdown: str) -> list[str]:
    """从 Markdown 表格里取出非空单元格（去掉分隔行与表格线）。"""
    cells: list[str] = []
    for line in markdown.split("\n"):
        if set(line.replace("|", "").strip()) <= {"-", " "}:
            continue  # 分隔行
        for cell in line.strip().strip("|").split("|"):
            value = cell.strip().replace("／", "|")
            if value:
                cells.append(value)
    return cells


def _is_subsequence(needle: str, haystack: str) -> bool:
    """`needle` 是否是 `haystack` 的子序列（逐字符、保持顺序、允许跳）。"""
    it = iter(haystack)
    return all(ch in it for ch in needle)


def _raw_text(path: Path) -> str:
    """不经我们的 loader，直接把原文取出来（作为清洗的对照基准）。"""
    if path.suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if path.suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        parts = [p.text or "" for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text or "" for cell in row.cells)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


# ============================================================ 3. 断句与拼接判据


@pytest.mark.parametrize(
    ("tail", "expected"),
    [
        ("这句话说完了。", False),
        ("This sentence ends.", False),
        ("问题解决了？", False),
        ("注意到这一点", True),  # 汉字结尾，很可能被切断
        ("the sentence was cut", True),  # 字母结尾
        ("前面有逗号，", True),  # 逗号结尾 = 没说完
        ("", False),
    ],
)
def test_looks_unfinished(tail: str, expected: bool) -> None:
    assert looks_unfinished(tail) is expected


def test_can_join_refuses_to_glue_body_onto_heading() -> None:
    """**实测踩到的坑**：标题行结尾没有标点，会被"没说完"判据误判。

    `第 1 章 · 第 1 页主题` 后面直接接上正文，读者看到的是标题和正文糊成一行 ——
    而这种错看起来像排版，不像 bug，最容易留在库里。
    """
    assert can_join("第 1 章 · 第 1 页主题", "本节是正文内容。") is False
    assert can_join("第六条 违约责任", "乙方逾期交付的。") is False


def test_can_join_glues_a_real_wrapped_sentence() -> None:
    assert can_join("甲方应在收到发票后三十个工作日内", "完成付款，逾期计息。") is True


def test_can_join_does_not_glue_two_finished_sentences() -> None:
    assert can_join("这句话说完了。", "下一句也完了。") is False


def test_merge_wrapped_lines_joins_pdf_line_breaks() -> None:
    """PDF 的换行是**排版行为**，不是段落边界 —— 不合并会让切分切断句子。"""
    merged = merge_wrapped_lines(
        [
            "甲方应在每一期交付物验收合格后十五个",
            "工作日内完成付款。",
            "以上为完整条款。",
        ]
    )
    assert merged[0].startswith("甲方应在每一期交付物验收合格后十五个工作日内完成付款。")
    assert len(merged) == 2


def test_merge_wrapped_lines_keeps_blank_line_as_paragraph_break() -> None:
    """空行是**真段落边界**，绝不能跨空行合并。"""
    merged = merge_wrapped_lines(["没说完的话", "", "另一段开始"])
    assert merged == ["没说完的话", "", "另一段开始"]


def test_merge_wrapped_lines_adds_space_between_english_words() -> None:
    """英文断行要补空格 —— 否则 `Machine` + `Learning` 粘成新词。"""
    merged = merge_wrapped_lines(["We use Machine", "Learning models."])
    assert merged[0] == "We use Machine Learning models."


# ============================================================ 4. 页眉页脚 / 页码行


def test_strip_repeated_header_footer_keeps_body() -> None:
    """**核心断言**：页眉页脚要删，正文一行不能少。

    这正是踩过的坑 —— 测试数据若每页正文相同，就无法区分
    "正常去页眉"和"丢了正文"。这里显式断言正文仍在。
    """
    pages = [
        ["机密文件", "正文甲 唯一", "正文乙 唯一", "第 1 页"],
        ["机密文件", "正文丙 唯一", "正文丁 唯一", "第 2 页"],
        ["机密文件", "正文戊 唯一", "正文己 唯一", "第 3 页"],
    ]
    cleaned = strip_repeated_lines(pages)

    for original, result in zip(pages, cleaned, strict=True):
        for line in original[1:-1]:
            assert line in result, f"正文行被误删：{line}，结果 {result}"
        assert "机密文件" not in result, "页眉应被剔除"


def test_strip_disabled_for_tiny_documents() -> None:
    """少于 3 页不做统计：2 页重复恰好 100%，会把正文当页脚删掉。"""
    pages = [["唯一内容 A"], ["唯一内容 A"]]
    assert strip_repeated_lines(pages) == pages


def test_page_number_lines_are_stripped_by_pattern() -> None:
    """页码每页都不同，频率统计抓不到 —— 必须用模式识别。

    但**只在首尾位置删**：正文里的孤立数字（金额）是正常内容。
    """
    assert strip_page_numbers(["第 3 页 / 共 50 页", "正文", "看重"]) == ["正文", "看重"]
    assert strip_page_numbers(["正文", "第 3 页"]) == ["正文"]
    assert strip_page_numbers(["金额 128000 元", "另一行"]) == ["金额 128000 元", "另一行"]


# ============================================================ 5. PDF：页码对得上


def test_pdf_page_numbers_are_sequential_and_correct() -> None:
    """**验收核心**：50 页 PDF，页码必须 1..50 连续且对得上。"""
    result = load_document(FIXTURES / "sample50.pdf")
    pages = [d.page for d in result.documents]
    assert pages == list(range(1, 51))


def test_pdf_page_content_matches_its_page_number() -> None:
    """每页内容里的"第 N 章"必须与该页 `page=N` 一致 —— 页码正确的实质证明。

    只断言"有 50 个 Document"是不够的：页码全错也能过。
    """
    result = load_document(FIXTURES / "sample50.pdf")
    for doc in result.documents:
        assert doc.page is not None
        assert f"第 {doc.page} 章" in doc.text, f"第 {doc.page} 页内容对不上：{doc.text[:40]}"
        assert doc.page in doc.pages, f"第 {doc.page} 页不在自己的 pages 里：{doc.pages}"


def test_pdf_header_and_page_number_are_removed() -> None:
    """页眉文字和页码行都不该出现在正文里（否则 50 页重复 50 次，白烧 token）。"""
    body = _body(load_document(FIXTURES / "sample50.pdf"))
    assert "严格保密" not in body, "页眉应被剔除"
    assert "共 50 页" not in body, "页码行应被剔除"


def test_pdf_body_lines_are_preserved() -> None:
    """**反向断言**：剔除要精准，正文一行都不能少。"""
    result = load_document(FIXTURES / "sample50.pdf")
    for doc in result.documents:
        lines = [ln for ln in doc.text.split("\n") if ln.strip()]
        assert len(lines) >= 3, f"第 {doc.page} 页正文行数不足：{lines}"


def test_pdf_without_header_also_works() -> None:
    """没有页眉页脚的 PDF 不能被误剥离。"""
    result = load_document(FIXTURES / "sample5_no_header.pdf")
    assert len(result.documents) == 5
    assert result.warnings == []
    for doc in result.documents:
        assert f"第 {doc.page} 章" in doc.text


# ============================================================ 6. PDF：跨页断句


def test_pdf_crosspage_sentence_is_joined() -> None:
    """页末没说完 → 与下一页首行拼接。**这是"页码会不会错位"的第二处关键**。"""
    result = load_document(FIXTURES / "sample_crosspage.pdf")
    first = result.documents[0]
    assert "三十个工作日内" in first.text
    # 拼接成功：同一段里既有"页末那半句"也有"下一页的开头"
    assert "三十个工作日内完成第 1 项付款" in first.text


def test_pdf_crosspage_paragraph_remembers_both_pages() -> None:
    """**核心**：拼接后的段落横跨两页，`pages` 必须记全。

    只记起始页是"半个真话"：用户点进 p1 找不到那句的后半截，
    点进 p2 又找不到前半截 —— 而这正是"页码错位"最容易被用户点穿的地方。
    """
    result = load_document(FIXTURES / "sample_crosspage.pdf")
    first = result.documents[0]
    assert first.page == 1, "起始页应该是 1"
    assert first.pages == (1, 2), f"跨页段落必须记全两页，实际 {first.pages}"


def test_pdf_crosspage_last_page_keeps_unfinished_tail() -> None:
    """最后一页没有下一页可拼，未完结的句子就保持原样（不丢内容）。"""
    result = load_document(FIXTURES / "sample_crosspage.pdf")
    last = result.documents[-1]
    assert last.text.rstrip().endswith("三十个工作日内")


# ============================================================ 7. PDF：表格


def test_pdf_table_becomes_markdown_block() -> None:
    """表格转 Markdown，且**独立成块**（`block="table"`）。"""
    result = load_document(FIXTURES / "sample_table.pdf")
    tables = [d for d in result.documents if d.is_table]
    assert tables, "没有识别出表格"
    table = tables[0]
    assert table.text.startswith("| 里程碑 | 时间 | 金额 |")
    assert "| --- |" in table.text
    assert "38400.00" in table.text
    assert table.page == 1


def test_pdf_table_text_is_not_duplicated_in_body_flow() -> None:
    """**关键**：表格内容不能同时出现在正文里。

    pypdf 抽表格是把单元格按阅读顺序铺平的。不摘掉的话，同一份内容会进库两次：
    一次是铺平的文字垃圾，一次是结构化表格 —— 检索时两条都命中，
    上下文里出现两份互相矛盾的版本。
    """
    result = load_document(FIXTURES / "sample_table.pdf")
    body = "\n".join(d.text for d in result.documents if not d.is_table)

    assert "第五条 付款计划" in body, "表格前的正文丢了"
    assert "第六条 违约责任" in body, "表格后的正文丢了"
    # 表格里的数据不该再以"每个单元格一行"的形式出现
    body_lines = [ln.strip() for ln in body.split("\n") if ln.strip()]
    for cell in ("里程碑", "38400.00", "二期交付"):
        assert cell not in body_lines, f"表格单元格 {cell!r} 仍留在正文流里"


def test_pdf_table_order_follows_page_order() -> None:
    """表格紧跟它所在的页，不前后颠倒（摘要时顺序是语义的一部分）。"""
    result = load_document(FIXTURES / "sample_table.pdf")
    blocks = [(d.page, d.block) for d in result.documents]
    assert blocks == [(1, "text"), (1, "table")]


def test_table_to_markdown_escapes_pipe() -> None:
    """单元格里的 `|` 会破坏表格语法，必须转义成别的字符。"""
    markdown = table_to_markdown([["a|b", "c"], ["d", "e"]])
    assert markdown.split("\n")[0] == "| a／b | c |"


def test_table_to_markdown_pads_ragged_rows() -> None:
    """行列不齐（合并单元格常见）时要补齐，否则 Markdown 表格渲染错乱。"""
    markdown = table_to_markdown([["a", "b", "c"], ["d"], ["e", "f"]])
    lines = markdown.split("\n")
    assert lines[0] == "| a | b | c |"
    assert lines[2] == "| d |  |  |"


def test_table_to_markdown_empty() -> None:
    assert table_to_markdown([]) == ""


# ============================================================ 8. PDF：扫描件


def test_scanned_pdf_records_warning_not_crash() -> None:
    """整份是"扫描件"（无文本层）：返回空结果 + 明确 warning，**不抛异常**。

    这是本格最重要的行为约定 —— 处理失败 ≠ 内容缺失。
    """
    result = load_document(FIXTURES / "sample_no_text.pdf")
    assert result.documents == []
    assert any("OCR" in w for w in result.warnings), result.warnings


def test_scan_detection_warns_on_low_char_pages(tmp_path: Path) -> None:
    """有字但**极少**的页（< 20 字）要单独告警：文本层可能残缺（部分扫描）。

    这类页比"完全没字"更危险：它有内容，看起来正常，只是信息大量缺失。
    不打 warning，用户永远不知道。

    **样本刻意避开 "Page 1" 这种写法**：那正好命中页码行正则，会被当成页码剔掉，
    于是这页变成"空页"，测到的是另一条分支（空页告警）而不是扫描页告警。
    测试样本本身踩中被测逻辑的边界，是最常见的"测了个寂寞"。
    """
    from reportlab.pdfgen import canvas

    path = tmp_path / "almost_blank.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(60, 780, "扫描残页")  # 4 个字，远少于 SCAN_PAGE_CHARS
    pdf.showPage()
    pdf.save()

    result = load_pdf(path)
    warnings = " ".join(result.warnings)
    assert str(SCAN_PAGE_CHARS) in warnings, result.warnings
    assert "OCR" in warnings, result.warnings


def test_short_but_normal_pdf_is_not_flagged(tmp_path: Path) -> None:
    """**反向断言**：正常短文档不能被误判成扫描页。

    `sample_short.pdf` 只有一行字，但它是完好的文本层 —— 不该有扫描告警。
    判据如果写成"少于 100 字就算扫描"，短通知/短合同全会被误报。
    """
    result = load_document(FIXTURES / "sample_short.pdf")
    assert len(result.documents) == 1
    assert not any("OCR" in w or "扫描" in w for w in result.warnings), result.warnings


def test_corrupt_pdf_raises_clear_error(tmp_path: Path) -> None:
    """坏文件是**真失败**，必须抛 —— 但错误信息要能看懂。"""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a real pdf at all")

    with pytest.raises(ValueError, match="无法打开 PDF"):
        load_document(path)


# ============================================================ 9. DOCX


def test_docx_extracts_paragraphs_and_tables() -> None:
    """**表格必须抽出来**：合同里的金额/时间表通常在表格里，只抽段落会丢核心信息。"""
    body = _body(load_document(FIXTURES / "sample.docx"))
    assert "技术服务合同" in body
    assert "合同编号：HT-2026-001" in body
    assert "一期交付" in body
    assert "38,400.00" in body


def test_docx_has_no_page_numbers() -> None:
    """DOCX 没有物理页概念，`page` 必须是 None 而不是硬编 1。

    硬编页码会让"第 3 页"变成谎言（Word 的页是渲染结果，文件里没存）。
    """
    result = load_document(FIXTURES / "sample.docx")
    assert all(d.page is None for d in result.documents)
    assert all(d.pages == () for d in result.documents)


def test_docx_heading_levels_go_into_metadata() -> None:
    """标题层级进 metadata：切分要靠它"遇到标题就断开"，引用要靠它显示章节名。"""
    result = load_document(FIXTURES / "sample.docx")
    headings = [d for d in result.documents if d.block == "heading"]
    assert headings, "没有识别出任何标题"
    levels = {d.text: d.metadata.get("heading_level") for d in headings}
    assert levels["技术服务合同"] == 1
    assert levels["第一条 服务范围"] == 2


def test_docx_table_is_a_separate_block() -> None:
    """表格必须是独立块（`block="table"`），不混进文字流。"""
    result = load_document(FIXTURES / "sample.docx")
    tables = [d for d in result.documents if d.is_table]
    assert tables, "没有识别出表格块"
    assert tables[0].text.startswith("| 里程碑 | 时间 | 金额 |")
    assert "| --- |" in tables[0].text


def test_docx_empty_document_warns() -> None:
    """空文档：返回空结果 + warning，不抛异常。"""
    result = load_document(FIXTURES / "empty.docx")
    assert result.documents == []
    assert result.warnings


@pytest.mark.parametrize(
    ("style", "expected"),
    [
        ("Heading 1", 1),
        ("Heading 3", 3),
        ("标题 1", 1),  # 中文版 Word 的样式名
        ("标题 2", 2),
        ("Title", 1),
        ("Normal", None),
        ("", None),
    ],
)
def test_docx_heading_style_parsing(style: str, expected: int | None) -> None:
    """**中文版 Word 的样式名是"标题 1"**：只匹配英文名会静默失效
    （所有标题都被当正文，且没有任何报错）。"""
    assert heading_level(style) == expected


# ============================================================ 10. Markdown


def test_markdown_splits_by_headings() -> None:
    """按标题切块：每个块是语义单元，块首即小标题。"""
    result = load_document(FIXTURES / "sample.md")
    assert len(result.documents) >= 4
    assert result.documents[0].text.startswith("# 工业 AI 平台设计说明")


def test_markdown_code_fence_does_not_split_on_hash() -> None:
    """代码块里的 `#` 是注释，不是标题 —— 必须在围栏内跳过标题识别。"""
    result = load_document(FIXTURES / "sample.md")
    code_docs = [d for d in result.documents if "```" in d.text]
    assert code_docs, "代码块没有被保留"
    assert "def handler" in code_docs[0].text


def test_markdown_table_is_separate_block() -> None:
    """Markdown 表格单独成块，与 PDF/DOCX 一致（同一套下游处理）。"""
    result = load_document(FIXTURES / "sample_table.md")
    tables = [d for d in result.documents if d.is_table]
    assert len(tables) == 1
    assert tables[0].text.startswith("| 里程碑 | 时间 | 金额 |")


def test_markdown_code_fence_pipe_is_not_table() -> None:
    """**代码块里的 `|` 是位运算，不是表格行**。

    不识别围栏的话，一段代码会被标成"表格"然后禁止切分 ——
    超大块进 embedding，表格语义还被错误挂到代码上。
    """
    result = load_document(FIXTURES / "sample_table.md")
    table_docs = [d for d in result.documents if d.is_table]
    assert len(table_docs) == 1, "代码块被误判成表格"
    assert all("flags = a | b" not in d.text for d in table_docs)


def test_markdown_heading_levels_go_into_metadata() -> None:
    result = load_document(FIXTURES / "sample_table.md")
    levels = [d.metadata.get("heading_level") for d in result.documents]
    assert 1 in levels
    assert 2 in levels


def test_markdown_block_with_body_is_not_marked_heading() -> None:
    """只有"整块就是一行标题"才标 `heading`。

    块里若还有正文，标成 heading 会让下游误以为整段都是标题 ——
    元数据说谎比没有元数据更糟（下游会据此做决策）。
    """
    result = load_document(FIXTURES / "sample.md")
    heading_blocks = [d for d in result.documents if d.block == "heading"]
    for doc in heading_blocks:
        assert len([ln for ln in doc.text.split("\n") if ln.strip()]) == 1


def test_markdown_utf8_and_gbk(tmp_path: Path) -> None:
    """Windows 上 GBK 编码的 .md 也要能读，且给出明确 warning。"""
    path = tmp_path / "gbk.md"
    path.write_bytes("# 标题\n\n正文内容。".encode("gbk"))

    result = load_document(path)
    assert result.documents
    assert any("GBK" in w for w in result.warnings), result.warnings


def test_markdown_empty_file_warns() -> None:
    result = load_document(FIXTURES / "empty.md")
    assert result.documents == []
    assert result.warnings


def test_markdown_whitespace_only_file_warns(tmp_path: Path) -> None:
    """ "只有空白"和"完全空"是两条不同的输入路径，都要给出同样的结论。

    **这个用例用 `tmp_path` 现场造，不进仓库**：只有空白的文件会被
    pre-commit 的 `trailing-whitespace` 钩子剥成 0 字节，
    于是"脚本产物"与"已提交的 fixture"漂移 —— 那种 diff 最让人困惑。
    """
    path = tmp_path / "whitespace.md"
    path.write_text("   \n\n  \n", encoding="utf-8")
    result = load_document(path)
    assert result.documents == []
    assert result.warnings


# ============================================================ 11. 三格式一致


def test_registry_covers_three_formats() -> None:
    assert set(SUPPORTED_EXTENSIONS) >= {".pdf", ".docx", ".md"}


def test_unsupported_extension_raises() -> None:
    with pytest.raises(ValueError, match="不支持的格式"):
        load_document(FIXTURES / "sample.txt")


def test_extension_is_case_insensitive(tmp_path: Path) -> None:
    """`.PDF` 大写也要认 —— Windows 上用户经常拿到大写后缀的文件。"""
    src = FIXTURES / "sample5_no_header.pdf"
    dst = tmp_path / "UPPER.PDF"
    dst.write_bytes(src.read_bytes())
    assert load_document(dst).documents


def test_missing_file_raises() -> None:
    with pytest.raises(ValueError, match="文件不存在"):
        load_document(FIXTURES / "nope.pdf")


def test_all_loaders_return_same_model() -> None:
    """**三格式统一输出的证明**：下游完全不需要知道源文件是什么格式。"""
    for name in ("sample50.pdf", "sample.docx", "sample.md"):
        result = load_document(FIXTURES / name)
        for doc in result.documents:
            assert isinstance(doc, Document)
            assert isinstance(doc.text, str) and doc.text
            # 三个字段三格式都有：page / pages / block
            assert doc.page is None or doc.page >= 1
            assert isinstance(doc.pages, tuple)
            assert doc.block in {"text", "table", "heading"}


def test_all_loaders_mark_tables_the_same_way() -> None:
    """**三格式一致的关键一条**：表格都用 `block="table"` + Markdown 文本。

    各 loader 自己发明格式，最终一定会有一处漏改 —— 到那时
    "为什么 PDF 的表格能答、DOCX 的答不了"要查半天。
    """
    for name in ("sample_table.pdf", "sample.docx", "sample_table.md"):
        tables = [d for d in load_document(FIXTURES / name).documents if d.is_table]
        assert tables, f"{name} 没有识别出表格块"
        assert tables[0].text.startswith("|"), f"{name} 的表格不是 Markdown"
        assert "| --- |" in tables[0].text, f"{name} 的表格没有分隔行"


def test_loader_registry_is_the_only_place_to_add_formats() -> None:
    """加格式只改 `LOADERS` 一处。这条测试守的是"分发集中"这个设计。"""
    assert set(LOADERS) == set(SUPPORTED_EXTENSIONS)
