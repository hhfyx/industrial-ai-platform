"""生成测试文档的 fixture 工厂。

**为什么要脚本化生成，而不是把二进制文件塞进仓库**：
- 二进制 PDF 在 code review 里完全不可读，没人能确认里面到底是什么内容
- 生成脚本本身是"测试意图的文档"：读者一眼看到"50 页，每页 3 段正文 + 页眉"，
  就知道这个 fixture 在测什么
- 需要不同形态的样本（有页眉 / 无页眉 / 扫描件）时，改参数即可

**关键设计：每页正文必须"每页都不同"。**
我第一版犯了相反的错误 —— 每页放了三行**完全相同**的句子。结果
`_strip_repeated_lines` 正确地把它识别成了"页脚"并删掉，
测试却看起来像"loader 丢内容"。

这是一个值得记住的教训：**如果测试数据的每页内容完全一样，
你就无法区分"loader 正常去页眉"和"loader 丢数据"。**
真实文档的正文绝不重复，所以 fixture 也必须做到。
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).resolve().parent

# 每页真正重复的部分（页眉 + 页码），这是要被 loader 剔除的
_HEADER = "星辰科技内部资料 · 严格保密"


def _register_cjk_font() -> str | None:
    """注册一个支持中文的字体。找不到就返回 None（退回 Helvetica，中文会变方块）。

    **必须在生成时就解决**：用 `reportlab` 默认字体画中文，抽出来全是乱码/方块，
    那后续所有"页码对不对、内容全不全"的断言都会被乱码干扰。
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("CJK", path))
                return "CJK"
            except Exception:  # noqa: BLE001, S112 - 换下一个候选
                continue
    return None


def make_pdf(path: Path, *, pages: int = 50, with_header: bool = True) -> Path:
    """生成一份多页 PDF，每页正文**各不相同**。

    结构（`with_header=True` 时）：
        [_HEADER]            ← 每页相同 → 应被删除
        [第 N 页 · 章节 N]   ← 每页不同 → **必须保留**
        [正文 3 行，含 N]    ← 每页不同 → **必须保留**
        [第 N 页 / 共 50 页] ← 页码行 → 应被删除
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)

    for index in range(1, pages + 1):
        y = 800
        if with_header:
            pdf.setFont(font or "Helvetica", 8) if font else pdf.setFont("Helvetica", 8)
            pdf.drawString(60, y, _HEADER)
            y -= 44

        pdf.setFont(font or "Helvetica", 15) if font else pdf.setFont("Helvetica", 13)
        pdf.drawString(60, y, f"第 {index} 章 · 第 {index} 页主题")
        y -= 28

        pdf.setFont(font or "Helvetica", 10) if font else pdf.setFont("Helvetica", 9)
        lines = [
            f"本节是第 {index} 页的正文，讨论第 {index} 项交付物的验收标准。",
            f"甲方应于第 {index} 次里程碑验收后十五个工作日内完成付款。",
            f"若第 {index} 期交付延期，乙方按合同总额的 {index}‰ 承担违约责任。",
        ]
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 18

        if with_header:
            pdf.setFont(font or "Helvetica", 8) if font else pdf.setFont("Helvetica", 8)
            pdf.drawString(60, 60, f"第 {index} 页 / 共 {pages} 页")

        pdf.showPage()

    pdf.save()
    return path


def make_docx(path: Path) -> Path:
    """生成一份 DOCX：包含标题、段落、表格。"""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("技术服务合同", level=1)
    doc.add_paragraph("甲方：杭州星辰科技有限公司")
    doc.add_paragraph("乙方：上海云脉信息技术有限公司")
    doc.add_paragraph("合同编号：HT-2026-001")

    doc.add_heading("第一条 服务范围", level=2)
    doc.add_paragraph("乙方为甲方提供工业设备预测性维护系统的开发与部署服务。")

    doc.add_heading("第二条 付款方式", level=2)
    doc.add_paragraph("甲方在验收后十五个工作日内支付合同金额的百分之九十。")

    table = doc.add_table(rows=3, cols=3)
    rows = [
        ("里程碑", "时间", "金额"),
        ("一期交付", "2026-11-30", "¥ 38,400.00"),
        ("二期交付", "2027-03-31", "¥ 89,600.00"),
    ]
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value

    doc.save(str(path))
    return path


def make_markdown(path: Path) -> Path:
    """生成一份 Markdown：多级标题 + 代码块（含 `#` 注释，测围栏识别）。"""
    content = """# 工业 AI 平台设计说明

本文档描述平台的总体架构与数据流。

## 一、整体架构

平台由接入层、调度层、模型层三部分组成。

### 1.1 接入层

接入层负责协议转换与限流，支持 HTTP 与 SSE。

```python
# 这段代码里的 # 不是标题，不该被切成新块
def handler(request):
    return {"ok": True}
```

## 二、数据流

原始文档经 loader 解析后，进入切分器，再按 token 预算分批送入模型。

## 三、部署

使用 Docker 多阶段构建，最终镜像约 157 MB。
"""
    path.write_text(content, encoding="utf-8")
    return path


def make_pdf_with_table(path: Path, *, with_header: bool = False) -> Path:
    """生成一份**带真实表格**的 PDF（有网格线，pdfplumber 才识别得出）。

    **为什么要网格线**：pdfplumber 的表格识别基于"线框/对齐的空隙"。
    只有文字对齐、没有线框时，它可能识别成一段普通文本 ——
    那样这个 fixture 就测不到"表格抽取"这条路径，测试看起来绿，
    实际什么都没验证。

    结构：
        第五条 付款计划        ← 正文前
        甲方按下列...          ← 正文
        [表格 3x3，带网格]     ← 应被单独抽成 block=table，且**不进文字流**
        第六条 违约责任        ← 正文后
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    use = font or "Helvetica"

    pdf.setFont(use, 12)
    pdf.drawString(60, 800, "第五条 付款计划")
    pdf.setFont(use, 10)
    pdf.drawString(60, 775, "甲方按下列里程碑分期支付合同价款。")

    rows = [
        ("里程碑", "时间", "金额"),
        ("一期交付", "2026-11-30", "38400.00"),
        ("二期交付", "2027-03-31", "89600.00"),
    ]
    x0, y_top, col_w, row_h = 60, 740, 150, 22
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            pdf.drawString(x0 + c * col_w + 4, y_top - r * row_h - 15, value)
    for r in range(len(rows) + 1):
        pdf.line(x0, y_top - r * row_h, x0 + 3 * col_w, y_top - r * row_h)
    for c in range(4):
        pdf.line(x0 + c * col_w, y_top, x0 + c * col_w, y_top - len(rows) * row_h)

    bottom = y_top - len(rows) * row_h - 40
    pdf.setFont(use, 12)
    pdf.drawString(60, bottom, "第六条 违约责任")
    pdf.setFont(use, 10)
    pdf.drawString(60, bottom - 25, "乙方逾期交付的，按日万分之五支付违约金。")

    pdf.showPage()
    pdf.save()
    return path


def make_pdf_crosspage(path: Path, *, pages: int = 3) -> Path:
    """生成一份**句子被页边界切断**的 PDF。

    刻意让每页最后一行没有句号，下一页第一行是续行：
        第 1 页末行: 甲方应在收到发票后三十个工作日内
        第 2 页首行: 完成付款，逾期按日万分之五计息。
    正确行为是**拼接**，而且拼接后的段落 `pages` 要同时记 (1, 2)。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    use = font or "Helvetica"

    for index in range(1, pages + 1):
        pdf.setFont(use, 11)
        # 页首：上一页那句话的**续行**（除第 1 页）
        y = 800
        if index > 1:
            pdf.drawString(60, y, f"完成第 {index - 1} 项付款，逾期按日万分之五计息。")
            y -= 30
        pdf.drawString(60, y, f"第 {index} 页正文起点，说明第 {index} 项交付内容。")
        y -= 24
        # 页末：一句话没说完（**故意不加句号**）
        pdf.drawString(60, y, f"甲方应在收到第 {index} 批发票后三十个工作日内")
        pdf.showPage()

    pdf.save()
    return path


def make_pdf_no_text(path: Path, *, pages: int = 3) -> Path:
    """生成一份**没有任何文字**的 PDF（模拟纯扫描件）。

    只画方块，不画文字 —— pypdf 抽出来是空串。用来测"扫描件必须打 warning"。
    """
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path))
    for _ in range(pages):
        pdf.rect(100, 100, 200, 200, fill=1)
        pdf.showPage()
    pdf.save()
    return path


def make_pdf_short(path: Path) -> Path:
    """生成一份**内容很少但正常**的 PDF（单页，几十字）—— 测短文档不崩。

    **字数刻意压在 20 字以上**：这个 fixture 要测的是"短文档能正常解析"，
    不是"短页被判为扫描页"。两件事混在一个 fixture 里，测试失败时
    你分不清是 loader 坏了还是判据触发了 —— 所以扫描页判据用
    `make_pdf_no_text` 与测试里的内联样本单独测。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont(font or "Helvetica", 12)
    pdf.drawString(60, 780, "简短说明：本合同一式两份，双方各执一份，自签字之日起生效。")
    pdf.showPage()
    pdf.save()
    return path


def make_pdf_wrapped(path: Path) -> Path:
    """生成一份**段落被硬换行断开**的 PDF（模拟一行放不下就换行）。

    正确行为是**合并成一段**（PDF 的换行是排版行为，不是段落边界）。
    测试要断言的正是：`甲方应在...十五个` + `工作日内完成付款。` 变成一句话，
    而不是两段。

    同时**打一个反例**在末尾：一行 legit 的段落结尾（有句号），
    它不该和下一行粘在一起。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    use = font or "Helvetica"
    pdf.setFont(use, 10)

    y = 780
    lines = [
        "甲方应在每一期交付物经双方书面确认验收合格后十五个",
        "工作日内完成付款，付款方式为银行转账，账户信息以合同附件为准。",
        "以上为完整条款，双方确认无异议。",
        "另：本合同附件与本合同具有同等法律效力。",
    ]
    for line in lines:
        pdf.drawString(60, y, line)
        y -= 20

    pdf.showPage()
    pdf.save()
    return path


def make_empty_docx(path: Path) -> Path:
    """生成一份**完全空**的 DOCX（没有任何段落和表格）。"""
    from docx import Document as DocxDocument

    DocxDocument().save(str(path))
    return path


def make_empty_markdown(path: Path) -> Path:
    """生成一份**完全空**的 Markdown（0 字节）。

    **不能写成"只有空白"**：pre-commit 的 `trailing-whitespace` 钩子会把
    空白行剥掉，于是仓库里的文件变成 0 字节，而生成脚本写出来的是 3 行空白 ——
    "脚本产物"和"已提交的 fixture"不一致。这种漂移会让人在改完脚本重跑之后
    看到一批莫名其妙的 diff。

    "有内容但全是空白"是另一个用例，在 `tests/test_loaders.py` 里用
    `tmp_path` 现场造 —— 那种文件本来就不该进仓库。
    """
    path.write_text("", encoding="utf-8")
    return path


def make_markdown_with_table(path: Path) -> Path:
    """生成一份带 Markdown 表格 + 代码块里含 `|` 的文件。

    **代码块里的 `|` 是关键断言点**：位运算 `a | b` 不能被误判成表格行。
    """
    content = """# 付款计划

## 一、里程碑

下表列出各期交付的时间与金额。

| 里程碑 | 时间 | 金额 |
| --- | --- | --- |
| 一期 | 2026-11-30 | 38400.00 |
| 二期 | 2027-03-31 | 89600.00 |

## 二、校验逻辑

```python
# 这行里的 | 是位运算，不是表格
flags = a | b
```
"""
    path.write_text(content, encoding="utf-8")
    return path


def main() -> int:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)

    targets = [
        make_pdf(FIXTURE_DIR / "sample50.pdf", pages=50, with_header=True),
        make_pdf(FIXTURE_DIR / "sample5_no_header.pdf", pages=5, with_header=False),
        make_pdf_with_table(FIXTURE_DIR / "sample_table.pdf"),
        make_pdf_crosspage(FIXTURE_DIR / "sample_crosspage.pdf", pages=3),
        make_pdf_no_text(FIXTURE_DIR / "sample_no_text.pdf", pages=3),
        make_pdf_short(FIXTURE_DIR / "sample_short.pdf"),
        make_pdf_wrapped(FIXTURE_DIR / "sample_wrapped.pdf"),
        make_docx(FIXTURE_DIR / "sample.docx"),
        make_empty_docx(FIXTURE_DIR / "empty.docx"),
        make_markdown(FIXTURE_DIR / "sample.md"),
        make_markdown_with_table(FIXTURE_DIR / "sample_table.md"),
        make_empty_markdown(FIXTURE_DIR / "empty.md"),
    ]

    for p in targets:
        print(f"  {p.name:26} {p.stat().st_size:>8} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
